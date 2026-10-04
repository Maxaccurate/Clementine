from pathlib import Path
import sys,json,subprocess,hashlib,zipfile,csv
from pptx import Presentation
from pptx.util import Inches
from docx import Document
from openpyxl import Workbook,load_workbook
from PIL import Image
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parent.parent;APP=ROOT/'outputs/DesktopDrop';EXE=APP/'DesktopDrop.exe';DIR=ROOT/'work/office-fixtures';DIR.mkdir(exist_ok=True)
sys.path.insert(0,str(APP/'backend'))
import worker,office
deck=Presentation()
for i in range(2):
    slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=f'Office test slide {i+1}';slide.placeholders[1].text=f'第 {i+1} 页中文内容\nA useful sentence on slide {i+1}'
deck.save(DIR/'slides.pptx')
word=Document();word.add_heading('Office conversion test',0);word.add_paragraph('第一页文字：中文和 English');table=word.add_table(rows=2,cols=2);table.cell(0,0).text='Item';table.cell(0,1).text='Amount';table.cell(1,0).text='Test';table.cell(1,1).text='123';word.add_page_break();word.add_paragraph('Second page selectable text');word.save(DIR/'document.docx')
book=Workbook();sheet=book.active;sheet.title='Sales';sheet.append(['Name','Amount']);sheet.append(['中文项目',123]);sheet.append(['Formula','=SUM(B2,7)']);sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.print_area='A1:B3';second=book.create_sheet('第二张表');second.append(['Second sheet',42]);second.print_area='A1:B1';book.save(DIR/'sheets.xlsx')
(DIR/'source.csv').write_text('名称,金额\n测试,15\nliteral,=1+2\n',encoding='utf-8-sig')
checks=[]
def run(name,source,fmt):
    original=hashlib.sha256(source.read_bytes()).hexdigest()
    proc=subprocess.run([str(EXE),'--convert',fmt,str(source)],capture_output=True,timeout=180,creationflags=0x08000000)
    result=json.loads((APP/'last-cli-result.json').read_text(encoding='utf8'));row=result['Files'][0]
    if proc.returncode or row['Error']:raise RuntimeError(row['Error'] or proc.stderr.decode('utf8','replace'))
    output=Path(row['Output']);assert hashlib.sha256(source.read_bytes()).hexdigest()==original
    if fmt=='pdf':assert len(PdfReader(output).pages)>=1
    elif fmt in ('png','jpg'):
        images=list(output.glob('*.'+fmt)) if output.is_dir() else [output];assert images
        for path in images:assert Image.open(path).width>100
        if source.suffix=='.pptx':assert len(images)==2
    elif fmt=='txt':
        if output.is_dir():assert list(output.glob('*.txt'))
        else:assert output.read_text(encoding='utf-8-sig').strip()
    elif fmt=='html':assert (output/'index.html').exists()
    elif fmt in ('pptx','docx','xlsx','odp','odt','ods'):assert zipfile.is_zipfile(output)
    elif fmt in ('ppt','doc','xls'):assert output.read_bytes()[:8]==bytes.fromhex('d0cf11e0a1b11ae1')
    elif fmt in ('csv','tsv'):assert output.is_dir() and len(list(output.iterdir()))==2
    elif fmt=='json':assert len(json.loads(output.read_text()))==2
    return output

for source,formats in [('slides.pptx',['pdf','png','jpg','txt','ppt','odp']),('document.docx',['pdf','png','jpg','txt','doc','rtf','odt','html']),('sheets.xlsx',['pdf','png','jpg','csv','tsv','json','txt','xls','ods'])]:
    for fmt in formats:
        name=source+'→'+fmt
        print('Testing',name,flush=True)
        try:output=run(name,DIR/source,fmt);checks.append({'test':name,'passed':True})
        except Exception as error:checks.append({'test':name,'passed':False,'error':str(error)});print('FAIL',str(error)[:400],flush=True)
        (ROOT/'outputs/office-conversion-tests.json').write_text(json.dumps({'checks':checks,'passed':sum(x['passed'] for x in checks),'failed':sum(not x['passed'] for x in checks)},ensure_ascii=False,indent=2),encoding='utf8')

for legacy,target in [('ppt','pptx'),('doc','docx'),('xls','xlsx')]:
    sources=sorted(DIR.glob('*-converted.'+legacy))
    if sources:
        try:run(legacy+'→'+target,sources[0],target);checks.append({'test':legacy+'→'+target,'passed':True})
        except Exception as error:checks.append({'test':legacy+'→'+target,'passed':False,'error':str(error)})
for action,paths,params in [('officePdf',[str(DIR/'slides.pptx')],{'pageOrder':'2'}),('officeMergePDF',[str(DIR/'slides.pptx'),str(DIR/'document.docx')],{}),('removeMetadata',[str(DIR/'document.docx')],{'remove':'true'})]:
    try:
        result=DIR/(action+'.json');worker.execute({'Paths':paths,'Action':action,'Parameters':params},result);row=json.loads(result.read_text())['Files'][0];assert not row['Error'],row['Error'];output=Path(row['Output'])
        if action=='officePdf':assert len(PdfReader(output).pages)==1
        if action=='officeMergePDF':assert len(PdfReader(output).pages)==4
        checks.append({'test':action,'passed':True})
    except Exception as error:checks.append({'test':action,'passed':False,'error':str(error)})
report={'passed':sum(x['passed'] for x in checks),'failed':sum(not x['passed'] for x in checks),'checks':checks};(ROOT/'outputs/office-conversion-tests.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps(report,ensure_ascii=False),flush=True)
sys.exit(bool(report['failed']))
