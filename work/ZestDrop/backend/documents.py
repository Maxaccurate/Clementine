from pathlib import Path
import io,json,re
from PIL import Image,ImageDraw,ImageFont
import pypdfium2 as pdfium
from pypdf import PdfReader,PdfWriter
from common import *

def reader(path,params):
    doc=PdfReader(path)
    if doc.is_encrypted and not doc.decrypt(params.get('password','')): raise ValueError(T('PDF 需要正确的密码'))
    return doc

def metadata(path):
    doc=PdfReader(path)
    return {'Pages':len(doc.pages),**{str(k).lstrip('/'):str(v) for k,v in (doc.metadata or {}).items()}}

def images_document(paths,fmt,params):
    images=[]
    for path in paths:
        image=open_image(path)
        if 'A' in image.getbands():
            white=Image.new('RGB',image.size,'white');white.paste(image,mask=image.getchannel('A'));image=white
        images.append(image.convert('RGB'))
    with output_file(paths[0],'combined' if len(paths)>1 else 'converted','.'+fmt) as state:
        if fmt=='pdf': images[0].save(state['temp'],format='PDF',save_all=True,append_images=images[1:],resolution=150)
        else:
            from docx import Document
            from docx.shared import Inches
            doc=Document()
            for i,image in enumerate(images):
                if i: doc.add_page_break()
                data=io.BytesIO(); image.save(data,format='PNG'); data.seek(0)
                doc.add_picture(data,width=Inches(6))
            doc.save(state['temp'])
    return state['output']

def render_page(page):
    # 300 DPI, reduced for oversized pages so the bitmap stays within the same 80 MP limit as images.
    width,height=page.get_size(); scale=min(300/72,.999*(80_000_000/max(1,width*height))**.5)
    return page.render(scale=scale).to_pil()

def render_pdf(path,fmt,params):
    doc=pdfium.PdfDocument(str(path),password=params.get('password') or None)
    if len(doc)==1:
        with output_file(path,'page-1','.'+fmt) as state:
            image=render_page(doc[0]); save_image(image,state['temp'],fmt,92)
    else:
        with output_folder(path,'pages-'+fmt) as state:
            for index in range(len(doc)):
                image=render_page(doc[index]); save_image(image,state['temp']/f'page-{index+1:04d}.{fmt}',fmt,92)
    doc.close(); return state['output']

def pdf_word(path,params):
    import pymupdf as fitz
    doc=fitz.open(path)
    if doc.needs_pass and not doc.authenticate(params.get('password','')): raise ValueError(T('PDF 密码错误'))
    with output_file(path,'converted','.docx') as state:
        if any(page.get_text().strip() for page in doc):
            from pdf2docx import Converter
            converter=Converter(str(path),password=params.get('password',''))
            try: converter.convert(str(state['temp']),multi_processing=False)
            finally: converter.close()
        else:
            from docx import Document
            from docx.shared import Inches
            word=Document()
            for i,page in enumerate(doc):
                if i: word.add_page_break()
                image=io.BytesIO(page.get_pixmap(matrix=fitz.Matrix(2,2)).tobytes('png')); word.add_picture(image,width=Inches(6))
            word.save(state['temp'])
    doc.close(); return state['output']

def text_image(text,width=1200):
    font=ImageFont.truetype(str(FONT),28)
    lines=[]
    for paragraph in text.splitlines() or ['']:
        current=''
        for char in paragraph:
            if font.getlength(current+char)>width-100:
                lines.append(current); current=char
            else: current+=char
        lines.append(current)
    height=max(180,100+len(lines)*42)
    if width*height>80_000_000: raise ValueError(T('文本过长，单张图片过大；请选择 PDF 输出'))
    image=Image.new('RGB',(width,height),'white'); draw=ImageDraw.Draw(image)
    for i,line in enumerate(lines): draw.text((50,50+i*42),line,fill='#111111',font=font)
    return image

def text_pdf(path,text,params):
    import pymupdf as fitz
    doc=fitz.open(); font=fitz.Font(fontfile=str(FONT))
    margin=48; lines=[]
    for paragraph in text.splitlines() or ['']:
        current=''
        for char in paragraph:
            if font.text_length(current+char,fontsize=12)>499: lines.append(current); current=char
            else: current+=char
        lines.append(current)
    for start in range(0,len(lines),44):
        page=doc.new_page(width=595,height=842); page.insert_font(fontname='DocumentFont',fontfile=str(FONT))
        for index,line in enumerate(lines[start:start+44]): page.insert_text((margin,margin+18+index*16),line,fontname='DocumentFont',fontsize=12)
    with output_file(path,'converted','.pdf') as state: doc.save(state['temp'],garbage=4,deflate=True)
    doc.close(); return state['output']

def timestamp(seconds,vtt=False):
    milliseconds=round(float(seconds)*1000); hours,rest=divmod(milliseconds,3600000); minutes,rest=divmod(rest,60000); seconds,millis=divmod(rest,1000)
    return f'{hours:02}:{minutes:02}:{seconds:02}'+('.' if vtt else ',')+f'{millis:03}'

def subtitles(path,fmt,params):
    text=Path(path).read_text(encoding='utf-8-sig'); source=Path(path).suffix.lower(); cues=[]
    if source=='.txt':
        span=number(params,'cueSeconds',4)
        for i,line in enumerate(x for x in text.splitlines() if x.strip()): cues.append((timestamp(i*span),timestamp((i+1)*span),line))
    else:
        for block in re.split(r'\n\s*\n',text.replace('\r','')):
            lines=block.splitlines()
            timed=next((i for i,x in enumerate(lines) if '-->' in x),None)
            if timed is not None:
                start,end=lines[timed].split('-->',1); end=end.strip().split()[0]
                cues.append((start.strip(),end,'\n'.join(lines[timed+1:])))
    if fmt=='txt': output='\n\n'.join(cue[2] for cue in cues)
    else:
        output='WEBVTT\n\n' if fmt=='vtt' else ''
        for i,(start,end,content) in enumerate(cues):
            separator='.' if fmt=='vtt' else ','
            normalize=lambda s: ('00:'+s if len(s.split(':'))==2 else s).replace(',',separator).replace('.',separator)
            output+=(str(i+1)+'\n' if fmt=='srt' else '')+normalize(start)+' --> '+normalize(end)+'\n'+content+'\n\n'
    with output_file(path,'converted','.'+fmt) as state: state['temp'].write_text(output,encoding='utf8')
    return state['output']

def convert(path,fmt,params):
    source=Path(path).suffix.lower()
    if source in ('.srt','.vtt') or fmt in ('srt','vtt'): return subtitles(path,fmt,params)
    if source=='.pdf':
        if fmt in ('jpg','png'): return render_pdf(path,fmt,params)
        if fmt=='docx': return pdf_word(path,params)
        if fmt=='txt':
            doc=reader(path,params); text='\n\n'.join(page.extract_text() or '' for page in doc.pages)
            if not text.strip(): raise ValueError(T('PDF 没有可选文字；扫描件需先识别文字'))
            with output_file(path,'converted','.txt') as state: state['temp'].write_text(text,encoding='utf8')
            return state['output']
    if source=='.txt':
        text=Path(path).read_text(encoding='utf-8-sig')
        if fmt=='pdf': return text_pdf(path,text,params)
        with output_file(path,'converted','.'+fmt) as state: save_image(text_image(text),state['temp'],fmt)
        return state['output']
    raise ValueError(T('不支持此文档转换'))

def merge(paths,params):
    writer=PdfWriter()
    for path in paths: writer.append(reader(path,params))
    with output_file(paths[0],'merged','.pdf') as state:
        with open(state['temp'],'wb') as stream: writer.write(stream)
    writer.close(); return state['output']

def tool(path,action,params):
    if action=='compress':
        import pymupdf as fitz
        doc=fitz.open(path)
        if doc.needs_pass and not doc.authenticate(params.get('password','')): raise ValueError(T('PDF 密码错误'))
        quality=int(number(params,'quality',75)); edge=int(number(params,'maxEdge',2000)); target=int(number(params,'targetKB',0)*1024)
        seen=set()
        for page in doc:
            for item in page.get_images(full=True):
                xref=item[0]
                if xref in seen: continue
                seen.add(xref)
                extracted=doc.extract_image(xref)
                try:
                    image=Image.open(io.BytesIO(extracted['image'])).convert('RGB'); image.thumbnail((edge,edge))
                    data=io.BytesIO(); image.save(data,format='JPEG',quality=quality)
                    if len(data.getvalue())<len(extracted['image']): page.replace_image(xref,stream=data.getvalue())
                except (OSError,ValueError): pass
        with output_file(path,'compressed','.pdf') as state:
            doc.save(state['temp'],garbage=4,deflate=True)
            if target and state['temp'].stat().st_size>target: raise ValueError(T('无法在当前设置下达到目标大小，请调低质量或缩小尺寸'))
        doc.close(); return state['output']
    doc=reader(path,params)
    if action=='splitPDF':
        group=int(number(params,'pagesPerFile',1))
        if group<1: raise ValueError(T('每份页数必须大于零'))
        with output_folder(path,'split') as state:
            for start in range(0,len(doc.pages),group):
                writer=PdfWriter()
                for page in doc.pages[start:start+group]: writer.add_page(page)
                with open(state['temp']/f'pages-{start+1:04d}-{min(len(doc.pages),start+group):04d}.pdf','wb') as stream: writer.write(stream)
                writer.close()
        return state['output']
    writer=PdfWriter()
    order=params.get('pageOrder','').strip()
    indexes=[int(x.strip())-1 for x in order.split(',')] if order else list(range(len(doc.pages)))
    rotations=json.loads(params.get('rotations','{}'))
    for index in indexes:
        if index<0 or index>=len(doc.pages): raise ValueError(T('页码超出文档范围'))
        page=writer.add_page(doc.pages[index])
        angle=int(rotations.get(str(index+1),number(params,'rotation',0)))
        if angle: page.rotate(angle)
    if action=='removeMetadata':
        if not truth(params.get('remove','true')): writer.add_metadata({'/'+str(k).lstrip('/'):str(v) for k,v in json.loads(params.get('metadata','{}')).items() if v is not None})
        if '/Metadata' in writer._root_object: del writer._root_object['/Metadata']
    else: writer.add_metadata({str(k):str(v) for k,v in (doc.metadata or {}).items()})
    with output_file(path,action,'.pdf') as state:
        with open(state['temp'],'wb') as stream: writer.write(stream)
    writer.close(); return state['output']
