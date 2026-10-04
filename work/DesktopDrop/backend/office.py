from pathlib import Path
import csv,io,json,tempfile,zipfile,xml.etree.ElementTree as ET
from common import *

PRESENTATIONS={'.ppt','.pptx','.pptm','.pps','.ppsx','.odp'}
WORDS={'.doc','.docx','.docm','.rtf','.odt'}
SHEETS={'.xls','.xlsx','.xlsm','.xlsb','.ods','.csv','.tsv'}
OFFICE_EXT=PRESENTATIONS|WORDS|SHEETS

def family(path):
    ext=Path(path).suffix.lower()
    return 'presentation' if ext in PRESENTATIONS else 'word' if ext in WORDS else 'sheet'

def native(path,fmt,target):
    process([BASE/'DesktopDrop.exe','--office-export',Path(path).resolve(),fmt,Path(target).resolve()],timeout=300)

def data_rows(path,folder):
    ext=Path(path).suffix.lower()
    if ext in ('.csv','.tsv'):
        with open(path,encoding='utf-8-sig',newline='') as stream:return [{'Name':Path(path).stem,'Rows':list(csv.reader(stream,delimiter='\t' if ext=='.tsv' else ','))}]
    if ext in ('.xlsx','.xlsm'):
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT,'Excel.Application\\CLSID'):pass
        except FileNotFoundError:
            from openpyxl import load_workbook
            workbook=load_workbook(path,data_only=True,read_only=True);result=[{'Name':sheet.title,'Rows':[list(row) for row in sheet.iter_rows(values_only=True)]} for sheet in workbook.worksheets];workbook.close();return result
    output=folder/'sheets.json';native(path,'sheet-data',output);return json.loads(output.read_text(encoding='utf8'))

def text_without_office(path):
    ext=Path(path).suffix.lower()
    if ext in ('.pptx','.pptm','.ppsx'):
        from pptx import Presentation
        document=Presentation(path);parts=[]
        for i,slide in enumerate(document.slides):
            parts.append(f'第 {i+1} 页')
            for shape in slide.shapes:
                if shape.has_text_frame:parts.append(shape.text)
                if shape.has_table:
                    parts.extend('\t'.join(cell.text for cell in row.cells) for row in shape.table.rows)
        return '\n\n'.join(parts)
    if ext in ('.docx','.docm'):
        from docx import Document
        document=Document(path);parts=[p.text for p in document.paragraphs]
        parts.extend('\t'.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return '\n'.join(parts)
    raise ValueError('此文本导出需要本机 Microsoft Office')

def pdf_file(path,folder):
    if Path(path).suffix.lower() in ('.csv','.tsv'):path=csv_workbook(path,folder)
    output=folder/'document.pdf';native(path,'pdf',output);return output

def csv_workbook(path,folder):
    from openpyxl import Workbook
    workbook=Workbook();worksheet=workbook.active
    for row in data_rows(path,folder)[0]['Rows']:worksheet.append(row)
    for row in worksheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
    intermediate=folder/'sheet.xlsx';workbook.save(intermediate);return intermediate

def convert(path,fmt,params):
    path=Path(path);kind=family(path)
    with tempfile.TemporaryDirectory(prefix='desktopdrop-office-') as temporary:
        folder=Path(temporary)
        if fmt in ('png','jpg'):
            import pypdfium2 as pdfium
            pdf=pdf_file(path,folder);document=pdfium.PdfDocument(pdf)
            if len(document)==1:
                with output_file(path,'converted','.'+fmt) as state:save_image(document[0].render(scale=300/72).to_pil(),state['temp'],fmt,92)
            else:
                with output_folder(path,'slides-'+fmt if kind=='presentation' else 'pages-'+fmt) as state:
                    for index in range(len(document)):save_image(document[index].render(scale=300/72).to_pil(),state['temp']/f'page-{index+1:04d}.{fmt}',fmt,92)
            document.close();return state['output']
        if fmt in ('csv','tsv','json','txt') and kind=='sheet':
            sheets=data_rows(path,folder)
            if fmt=='json':
                with output_file(path,'converted','.json') as state:state['temp'].write_text(json.dumps(sheets,ensure_ascii=False,indent=2),encoding='utf8')
            elif len(sheets)==1:
                with output_file(path,'converted','.'+fmt) as state:
                    with open(state['temp'],'w',encoding='utf-8-sig',newline='') as stream:csv.writer(stream,delimiter=',' if fmt=='csv' else '\t').writerows(sheets[0]['Rows'])
            else:
                with output_folder(path,'sheets-'+fmt) as state:
                    for index,sheet in enumerate(sheets):
                        name=''.join('_' if c in '<>:"/\\|?*' else c for c in sheet['Name']).strip('. ') or f'Sheet{index+1}'
                        with open(state['temp']/f'{index+1:02d}-{name}.{fmt}','w',encoding='utf-8-sig',newline='') as stream:csv.writer(stream,delimiter=',' if fmt=='csv' else '\t').writerows(sheet['Rows'])
            return state['output']
        if fmt=='txt':
            with output_file(path,'converted','.txt') as state:
                if path.suffix.lower() in ('.pptx','.pptm','.ppsx','.docx','.docm'):state['temp'].write_text(text_without_office(path),encoding='utf-8-sig')
                else:native(path,'txt',state['temp'])
            return state['output']
        if fmt=='html':
            with output_folder(path,'html') as state:native(path,'html',state['temp']/'index.html')
            return state['output']
        if path.suffix.lower() in ('.csv','.tsv') and fmt in ('xlsx','xls','ods','pdf'):
            intermediate=csv_workbook(path,folder)
            with output_file(path,'converted','.'+fmt) as state:
                if fmt=='xlsx':shutil.copyfile(intermediate,state['temp'])
                else:native(intermediate,fmt,state['temp'])
            return state['output']
        with output_file(path,'converted','.'+fmt) as state:native(path,fmt,state['temp'])
        return state['output']

def metadata(path):
    if not zipfile.is_zipfile(path):return {'Format':Path(path).suffix.upper(),'Details':'此旧格式的元数据编辑暂不支持'}
    result={}
    with zipfile.ZipFile(path) as package:
        for name in ('docProps/core.xml','docProps/app.xml'):
            if name in package.namelist():
                for node in ET.fromstring(package.read(name)):
                    if node.text:result[node.tag.split('}')[-1]]=node.text
    return result

def edit_metadata(path,params):
    if not zipfile.is_zipfile(path):raise ValueError('元数据编辑支持现代 Office 格式，请先转换为 PPTX/DOCX/XLSX')
    values=json.loads(params.get('metadata','{}'));remove=truth(params.get('remove','true'))
    with output_file(path,'metadata',Path(path).suffix.lower()) as state:
        with zipfile.ZipFile(path) as source,zipfile.ZipFile(state['temp'],'w',zipfile.ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data=source.read(entry.filename)
                if entry.filename in ('docProps/core.xml','docProps/app.xml'):
                    root=ET.fromstring(data)
                    for node in list(root):
                        key=node.tag.split('}')[-1]
                        if remove and key not in ('Application','AppVersion','Slides','Words','Pages','Characters','Paragraphs','ScaleCrop','SharedDoc','LinksUpToDate','HyperlinksChanged'):node.text=''
                        elif not remove and key in values:node.text='' if values[key] is None else str(values[key])
                    data=ET.tostring(root,encoding='utf8',xml_declaration=True)
                if remove and entry.filename=='docProps/custom.xml':
                    root=ET.fromstring(data)
                    for node in list(root):root.remove(node)
                    data=ET.tostring(root,encoding='utf8',xml_declaration=True)
                target.writestr(entry,data)
    return state['output']

def inspect(path,folder):
    import pypdfium2 as pdfium
    folder=Path(folder);pdf=pdf_file(path,folder);document=pdfium.PdfDocument(pdf);preview=folder/'preview.png';document[0].render(scale=1.3).to_pil().save(preview)
    value={'Kind':'office','Preview':str(preview),'Width':0,'Height':0,'Duration':0,'FrameRate':0,'Pages':len(document),'Metadata':metadata(path)};document.close();return value

def selected_pdf(path,params):
    from pypdf import PdfReader,PdfWriter
    with tempfile.TemporaryDirectory(prefix='desktopdrop-office-') as temporary:
        source=pdf_file(path,Path(temporary));reader=PdfReader(source);order=params.get('pageOrder','').strip();indices=[int(v.strip())-1 for v in order.split(',')] if order else list(range(len(reader.pages)));writer=PdfWriter()
        for index in indices:
            if index<0 or index>=len(reader.pages):raise ValueError('页码超出文档范围')
            writer.add_page(reader.pages[index])
        with output_file(path,'selected','.pdf') as state:
            with open(state['temp'],'wb') as stream:writer.write(stream)
        writer.close();return state['output']

def merge(paths,params):
    from pypdf import PdfReader,PdfWriter
    writer=PdfWriter()
    with tempfile.TemporaryDirectory(prefix='desktopdrop-office-merge-') as temporary:
        for index,path in enumerate(paths):
            folder=Path(temporary)/str(index);folder.mkdir();pdf=pdf_file(path,folder);writer.append(PdfReader(pdf))
        with output_file(paths[0],'merged','.pdf') as state:
            with open(state['temp'],'wb') as stream:writer.write(stream)
    writer.close();return state['output']
